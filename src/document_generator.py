import os
import html

import networkx as nx
import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt

from PIL import Image as PILImage

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle
)
from reportlab.lib.units import mm

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image
)

from pptx import Presentation
from pptx.util import Pt, Inches


class DocumentGenerator:

    def __init__(self):

        base = getSampleStyleSheet()

        self.styles = {

            "Title": ParagraphStyle(
                "Title",
                parent=base["Title"],
                fontSize=24,
                leading=30,
                alignment=TA_CENTER,
                spaceAfter=8
            ),

            "Subtitle": ParagraphStyle(
                "Subtitle",
                parent=base["Normal"],
                fontSize=10,
                leading=14,
                alignment=TA_CENTER,
                textColor=colors.grey
            ),

            "Section": ParagraphStyle(
                "Section",
                parent=base["Heading2"],
                fontSize=16,
                leading=20,
                spaceBefore=14,
                spaceAfter=8
            ),

            "Body": ParagraphStyle(
                "Body",
                parent=base["BodyText"],
                fontSize=11,
                leading=16,
                spaceAfter=5
            ),

            "Example": ParagraphStyle(
                "Example",
                parent=base["BodyText"],
                fontSize=10.5,
                leading=15
            ),

            "Equation": ParagraphStyle(
                "Equation",
                parent=base["BodyText"],
                fontSize=12,
                leading=18,
                alignment=TA_CENTER
            ),

            "Small": ParagraphStyle(
                "Small",
                parent=base["BodyText"],
                fontSize=9,
                leading=12,
                textColor=colors.grey
            )
        }

    # --------------------------------------------------
    # SAFE TEXT
    # --------------------------------------------------

    def safe(self, text):

        if text is None:
            return ""

        return html.escape(
            str(text)
        )

    # --------------------------------------------------
    # EXAMPLE
    # --------------------------------------------------

    def add_example(
        self,
        story,
        text
    ):

        table = Table(
            [
                [
                    Paragraph(
                        f"<b>Example</b><br/>"
                        f"{self.safe(text)}",
                        self.styles["Example"]
                    )
                ]
            ]
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, -1),
                        colors.HexColor("#EEF5FF")
                    ),
                    (
                        "BOX",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.HexColor("#AAB8D4")
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        10
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        10
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        8
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        8
                    )
                ]
            )
        )

        story.append(table)
        story.append(
            Spacer(1, 8)
        )

    # --------------------------------------------------
    # ORIGINAL BOARD IMAGE
    # --------------------------------------------------

    def add_board_image(
        self,
        story,
        image_path
    ):

        if not os.path.exists(image_path):
            return

        img = PILImage.open(
            image_path
        )

        width, height = img.size

        max_width = 170 * mm

        scale = max_width / width

        story.append(
            Image(
                image_path,
                width=max_width,
                height=height * scale
            )
        )

        story.append(
            Spacer(1, 10)
        )

    # --------------------------------------------------
    # DRAW GRAPH
    # --------------------------------------------------

    def draw_graph(
        self,
        diagram,
        output_dir,
        index
    ):

        vertices = diagram.get(
            "vertices",
            []
        )

        edges = diagram.get(
            "edges",
            []
        )

        if not vertices:
            return None

        graph = nx.DiGraph() if diagram.get(
            "directed",
            False
        ) else nx.Graph()

        graph.add_nodes_from(
            vertices
        )

        weighted = False

        for edge in edges:

            if not isinstance(
                edge,
                dict
            ):
                continue

            source = edge.get(
                "from"
            )

            target = edge.get(
                "to"
            )

            if source is None or target is None:
                continue

            weight = edge.get(
                "weight"
            )

            graph.add_edge(
                source,
                target,
                weight=weight
            )

            if weight is not None:
                weighted = True

        # Automatically arrange the graph.
        position = nx.spring_layout(
            graph,
            seed=42
        )

        fig, ax = plt.subplots(
            figsize=(8, 5.5)
        )

        nx.draw_networkx_nodes(
            graph,
            position,
            node_size=1800,
            node_color="white",
            edgecolors="black",
            linewidths=1.5,
            ax=ax
        )

        nx.draw_networkx_labels(
            graph,
            position,
            font_size=12,
            font_weight="bold",
            ax=ax
        )

        nx.draw_networkx_edges(
            graph,
            position,
            arrows=diagram.get(
                "directed",
                False
            ),
            width=1.8,
            edge_color="black",
            ax=ax
        )

        if weighted:

            edge_labels = {}

            for source, target, data in graph.edges(
                data=True
            ):

                weight = data.get(
                    "weight"
                )

                if weight is not None:

                    edge_labels[
                        (source, target)
                    ] = str(weight)

            nx.draw_networkx_edge_labels(
                graph,
                position,
                edge_labels=edge_labels,
                font_size=11,
                ax=ax
            )

        title = diagram.get(
            "title",
            "Graph"
        )

        ax.set_title(
            title,
            fontsize=16,
            pad=15
        )

        ax.axis("off")

        os.makedirs(
            output_dir,
            exist_ok=True
        )

        output_path = os.path.join(
            output_dir,
            f"generated_graph_{index}.png"
        )

        plt.tight_layout()

        plt.savefig(
            output_path,
            dpi=200,
            bbox_inches="tight"
        )

        plt.close(fig)

        return output_path

    # --------------------------------------------------
    # DRAW HIERARCHY
    # --------------------------------------------------

    def draw_hierarchy(
        self,
        diagram,
        output_dir,
        index
    ):

        root = diagram.get(
            "root"
        )

        if not root:
            return None

        graph = nx.DiGraph()

        graph.add_node(
            root
        )

        def add_children(
            parent,
            children
        ):

            if not isinstance(
                children,
                list
            ):
                return

            for child in children:

                if not isinstance(
                    child,
                    dict
                ):
                    continue

                label = child.get(
                    "label"
                )

                if not label:
                    continue

                graph.add_edge(
                    parent,
                    label
                )

                add_children(
                    label,
                    child.get(
                        "children",
                        []
                    )
                )

        add_children(
            root,
            diagram.get(
                "children",
                []
            )
        )

        # Calculate tree levels.
        levels = {}

        levels[root] = 0

        queue = [root]

        while queue:

            current = queue.pop(0)

            current_level = levels[current]

            for child in graph.successors(
                current
            ):

                if child not in levels:

                    levels[child] = (
                        current_level + 1
                    )

                    queue.append(
                        child
                    )

        # Arrange nodes by hierarchy level.
        position = {}

        grouped = {}

        for node, level in levels.items():

            grouped.setdefault(
                level,
                []
            ).append(node)

        for level, nodes in grouped.items():

            count = len(nodes)

            for i, node in enumerate(nodes):

                x = (
                    i
                    - (count - 1) / 2
                )

                position[node] = (
                    x,
                    -level
                )

        fig, ax = plt.subplots(
            figsize=(8, 5)
        )

        nx.draw_networkx_nodes(
            graph,
            position,
            node_size=2200,
            node_color="white",
            edgecolors="black",
            linewidths=1.5,
            ax=ax
        )

        nx.draw_networkx_labels(
            graph,
            position,
            font_size=10,
            ax=ax
        )

        nx.draw_networkx_edges(
            graph,
            position,
            arrows=True,
            arrowsize=18,
            width=1.5,
            ax=ax
        )

        title = diagram.get(
            "title",
            root
        )

        ax.set_title(
            title,
            fontsize=16,
            pad=15
        )

        ax.axis("off")

        os.makedirs(
            output_dir,
            exist_ok=True
        )

        output_path = os.path.join(
            output_dir,
            f"generated_hierarchy_{index}.png"
        )

        plt.tight_layout()

        plt.savefig(
            output_path,
            dpi=200,
            bbox_inches="tight"
        )

        plt.close(fig)

        return output_path

    # --------------------------------------------------
    # ADD GENERATED DIAGRAM TO PDF
    # --------------------------------------------------

    def add_diagram_image(
        self,
        story,
        image_path
    ):

        if not image_path:
            return

        if not os.path.exists(
            image_path
        ):
            return

        img = PILImage.open(
            image_path
        )

        width, height = img.size

        max_width = 165 * mm
        max_height = 110 * mm

        scale = min(
            max_width / width,
            max_height / height
        )

        story.append(
            Image(
                image_path,
                width=width * scale,
                height=height * scale
            )
        )

        story.append(
            Spacer(1, 12)
        )

    # --------------------------------------------------
    # GENERATE PDF
    # --------------------------------------------------

    def generate_pdf(
        self,
        content,
        output_path,
        page_images=None
    ):

        os.makedirs(
            os.path.dirname(
                output_path
            ),
            exist_ok=True
        )

        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm
        )

        story = []

        # --------------------------------------------------
        # TITLE
        # --------------------------------------------------

        heading = content.get(
            "heading",
            "Class Notes"
        )

        story.append(
            Paragraph(
                self.safe(heading),
                self.styles["Title"]
            )
        )

        story.append(
            Paragraph(
                "Digitized Whiteboard Notes",
                self.styles["Subtitle"]
            )
        )

        story.append(
            Spacer(1, 15)
        )

        # --------------------------------------------------
        # SECTIONS
        # --------------------------------------------------

        sections = content.get(
            "sections",
            []
        )

        for section in sections:

            if not isinstance(
                section,
                dict
            ):
                continue

            title = section.get(
                "title",
                ""
            )

            if title:

                story.append(
                    Paragraph(
                        self.safe(title),
                        self.styles["Section"]
                    )
                )

            section_content = section.get(
                "content",
                []
            )

            for item in section_content:

                if not isinstance(
                    item,
                    dict
                ):
                    continue

                item_type = item.get(
                    "type",
                    "text"
                )

                text = item.get(
                    "text",
                    ""
                )

                if not text:
                    continue

                if item_type == "example":

                    self.add_example(
                        story,
                        text
                    )

                elif item_type == "equation":

                    story.append(
                        Paragraph(
                            self.safe(text),
                            self.styles["Equation"]
                        )
                    )

                    story.append(
                        Spacer(1, 6)
                    )

                elif item_type == "bullet":

                    story.append(
                        Paragraph(
                            f"• {self.safe(text)}",
                            self.styles["Body"]
                        )
                    )

                elif item_type == "definition":

                    story.append(
                        Paragraph(
                            f"<b>Definition:</b> "
                            f"{self.safe(text)}",
                            self.styles["Body"]
                        )
                    )

                elif item_type == "question":

                    story.append(
                        Paragraph(
                            f"<b>Question:</b> "
                            f"{self.safe(text)}",
                            self.styles["Body"]
                        )
                    )

                else:

                    story.append(
                        Paragraph(
                            self.safe(text),
                            self.styles["Body"]
                        )
                    )

        # --------------------------------------------------
        # DIAGRAMS
        # --------------------------------------------------

        diagrams = content.get(
            "diagrams",
            []
        )

        if diagrams:

            story.append(
                Paragraph(
                    "Diagrams",
                    self.styles["Section"]
                )
            )

            diagram_dir = os.path.join(
                os.path.dirname(output_path),
                "generated_diagrams"
            )

            os.makedirs(
                diagram_dir,
                exist_ok=True
            )

            for index, diagram in enumerate(
                diagrams,
                start=1
            ):

                if not isinstance(
                    diagram,
                    dict
                ):
                    continue

                diagram_type = diagram.get(
                    "type"
                )

                generated_image = None

                if diagram_type == "graph":

                    generated_image = (
                        self.draw_graph(
                            diagram,
                            diagram_dir,
                            index
                        )
                    )

                elif diagram_type == "hierarchy":

                    generated_image = (
                        self.draw_hierarchy(
                            diagram,
                            diagram_dir,
                            index
                        )
                    )

                if generated_image:

                    self.add_diagram_image(
                        story,
                        generated_image
                    )

        # --------------------------------------------------
        # KEY TAKEAWAYS
        # --------------------------------------------------

        key_takeaways = content.get(
            "key_takeaways",
            []
        )

        if key_takeaways:

            story.append(
                Paragraph(
                    "Key Takeaways",
                    self.styles["Section"]
                )
            )

            for item in key_takeaways:

                story.append(
                    Paragraph(
                        f"• {self.safe(item)}",
                        self.styles["Body"]
                    )
                )

        # --------------------------------------------------
        # QUESTIONS
        # --------------------------------------------------

        questions = content.get(
            "questions",
            []
        )

        if questions:

            story.append(
                Paragraph(
                    "Questions",
                    self.styles["Section"]
                )
            )

            for index, question in enumerate(
                questions,
                start=1
            ):

                story.append(
                    Paragraph(
                        f"{index}. "
                        f"{self.safe(question)}",
                        self.styles["Body"]
                    )
                )

        # --------------------------------------------------
        # ORIGINAL BOARD IMAGES
        # --------------------------------------------------

        if page_images:

            story.append(
                Paragraph(
                    "Original Whiteboard Pages",
                    self.styles["Section"]
                )
            )

            story.append(
                Paragraph(
                    "Original board captures are included "
                    "to preserve the teacher's handwriting, "
                    "diagrams and visual layout.",
                    self.styles["Small"]
                )
            )

            story.append(
                Spacer(1, 8)
            )

            for index, image_path in enumerate(
                page_images,
                start=1
            ):

                if not os.path.exists(
                    image_path
                ):
                    continue

                story.append(
                    Paragraph(
                        f"<b>Board Page {index}</b>",
                        self.styles["Body"]
                    )
                )

                self.add_board_image(
                    story,
                    image_path
                )

        # --------------------------------------------------
        # BUILD
        # --------------------------------------------------

        doc.build(
            story
        )

        print(
            f"PDF saved at {output_path}"
        )

    # --------------------------------------------------
    # GENERATE PPT
    # --------------------------------------------------

    def generate_ppt(
        self,
        content,
        output_path,
        page_images=None
    ):

        prs = Presentation()

        # Title slide
        slide = prs.slides.add_slide(
            prs.slide_layouts[0]
        )

        slide.shapes.title.text = content.get(
            "heading",
            "Class Notes"
        )

        slide.placeholders[1].text = (
            "Digitized Whiteboard Notes"
        )

        # --------------------------------------------------
        # TEXT SLIDES
        # --------------------------------------------------

        def add_slide(
            title,
            lines
        ):

            if not lines:
                return

            slide = prs.slides.add_slide(
                prs.slide_layouts[1]
            )

            slide.shapes.title.text = title

            text_frame = (
                slide.placeholders[1]
                .text_frame
            )

            text_frame.clear()

            for index, line in enumerate(
                lines
            ):

                paragraph = (
                    text_frame.paragraphs[0]
                    if index == 0
                    else text_frame.add_paragraph()
                )

                paragraph.text = str(line)

                paragraph.level = 0

                paragraph.font.size = Pt(20)

        # Convert sections into slides.
        for section in content.get(
            "sections",
            []
        ):

            if not isinstance(
                section,
                dict
            ):
                continue

            title = section.get(
                "title",
                "Notes"
            )

            lines = []

            for item in section.get(
                "content",
                []
            ):

                if isinstance(
                    item,
                    dict
                ):

                    text = item.get(
                        "text"
                    )

                    if text:
                        lines.append(
                            str(text)
                        )

            add_slide(
                title,
                lines
            )

        # --------------------------------------------------
        # ORIGINAL BOARD IMAGES
        # --------------------------------------------------

        if page_images:

            blank = prs.slide_layouts[5]

            for index, image_path in enumerate(
                page_images,
                start=1
            ):

                if not os.path.exists(
                    image_path
                ):
                    continue

                slide = prs.slides.add_slide(
                    blank
                )

                title_box = (
                    slide.shapes.add_textbox(
                        Inches(0.5),
                        Inches(0.2),
                        Inches(8),
                        Inches(0.4)
                    )
                )

                title_box.text_frame.text = (
                    f"Original Whiteboard Page {index}"
                )

                slide.shapes.add_picture(
                    image_path,
                    Inches(0.4),
                    Inches(0.7),
                    width=Inches(8.3)
                )

        prs.save(
            output_path
        )

        print(
            f"PPT saved at {output_path}"
        )